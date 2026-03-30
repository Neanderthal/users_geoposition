import threading
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app.config import DATA_DIR

_lock = threading.Lock()

HEADERS = ["timestamp", "latitude", "longitude"]


def _ensure_workbook(path: Path) -> Workbook:
    if path.exists():
        return load_workbook(path)
    wb = Workbook()
    ws = wb.active
    ws.title = "locations"
    ws.append(HEADERS)
    return wb


def save_location(campaign_id: str, latitude: float, longitude: float) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    path = DATA_DIR / f"{campaign_id}.xlsx"

    with _lock:
        wb = _ensure_workbook(path)
        ws = wb.active
        ws.append([
            datetime.now(timezone.utc).isoformat(),
            latitude,
            longitude,
        ])
        wb.save(path)
